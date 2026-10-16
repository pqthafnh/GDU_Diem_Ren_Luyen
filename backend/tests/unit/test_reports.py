"""Unit tests for safe XLSX report generation."""

from openpyxl import load_workbook

from app.modules.reports.service import build_event_workbook, sanitize_excel_value


def test_excel_values_escape_formula_prefixes() -> None:
    """Spreadsheet formula prefixes are stored as literal text."""
    assert sanitize_excel_value("=SUM(A1)") == "'=SUM(A1)"
    assert sanitize_excel_value("normal") == "normal"


def test_event_workbook_contains_required_headers_and_rows() -> None:
    """The event workbook has Vietnamese headers and frozen/filterable rows."""
    workbook = load_workbook(
        build_event_workbook(
            [(1, "2024001", "Nguyễn Văn", "An", "Nam", None, "2401", "CNTT", "Hội thảo", 15)]
        )
    )
    sheet = workbook["Bao cao diem danh"]
    assert [cell.value for cell in sheet[1]] == [
        "STT",
        "MSSV",
        "Họ",
        "Tên",
        "Giới tính",
        "Ngày sinh",
        "Lớp",
        "Ngành",
        "Tên hoạt động",
        "Tổng",
    ]
    assert sheet.max_row == 2
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:J2"
