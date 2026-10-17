"""Unit tests for administrator student import validation."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.modules.admin_students.router import _excel_staff, _excel_students
from app.modules.admin_students.schemas import StudentCreate


def test_student_create_trims_identity_fields() -> None:
    """Student identity values are normalized before insertion."""
    student = StudentCreate(
        studentId=" 2025001 ",
        fullName=" Nguyễn Văn An ",
        email="an@gdu.edu.vn",
        initialPassword="Initial@123",
        birthDate=date(2007, 1, 1),
        enrollmentYear=2025,
    )
    assert student.student_id == "2025001"
    assert student.full_name == "Nguyễn Văn An"


def test_excel_import_reads_required_columns() -> None:
    """The importer maps the documented XLSX headers to validated students."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    headers = [
        "studentId",
        "fullName",
        "email",
        "initialPassword",
        "gender",
        "birthDate",
        "className",
        "major",
        "phoneNumber",
        "enrollmentYear",
    ]
    sheet.append(headers)
    sheet.append(
        [
            "2025002",
            "Trần Thị Bình",
            "binh@gdu.edu.vn",
            "Initial@123",
            "Nữ",
            "2007-02-02",
            "2501CNTT01",
            "Công nghệ thông tin",
            "0900000000",
            2025,
        ]
    )
    output = BytesIO()
    workbook.save(output)

    students = _excel_students(output.getvalue())

    assert len(students) == 1
    assert students[0].student_id == "2025002"
    assert students[0].enrollment_year == 2025


def test_excel_import_rejects_missing_columns() -> None:
    """The importer rejects files that cannot safely create complete accounts."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["studentId", "fullName"])
    output = BytesIO()
    workbook.save(output)

    with pytest.raises(ValueError, match="Missing Excel columns"):
        _excel_students(output.getvalue())


def test_excel_staff_import_reads_role() -> None:
    """The staff importer accepts only the ADMIN/MOD account contract."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(
        ["fullName", "email", "initialPassword", "role", "gender", "birthDate", "phoneNumber"]
    )
    sheet.append(
        ["Trần Văn Mod", "mod2@gdu.edu.vn", "Initial@123", "MOD", "Nam", "1988-01-01", "0900000001"]
    )
    output = BytesIO()
    workbook.save(output)

    staff = _excel_staff(output.getvalue())

    assert len(staff) == 1
    assert staff[0].role == "MOD"
