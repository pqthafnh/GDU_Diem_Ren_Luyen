"""End-to-end report and master matrix export verification."""

from io import BytesIO
from uuid import UUID

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook

from app.config.settings import get_settings
from app.database.pool import Database
from app.shared.security import create_access_token


def _headers(user_id: UUID, role: str) -> dict[str, str]:
    """Build a bearer header for a report actor."""
    settings = get_settings()
    token = create_access_token(
        user_id,
        role,
        f"{user_id}@test.gdu.edu.vn",
        settings.jwt_secret,
        settings.jwt_issuer,
        settings.jwt_audience,
        60,
    )
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.asyncio
async def test_event_report_and_master_matrix_exports_are_xlsx(
    client: AsyncClient,
    database: Database,
    standard_event: UUID,
    seeded_users: dict[str, UUID],
) -> None:
    """Export both report surfaces and verify workbook structure."""
    student_headers = _headers(seeded_users["STUDENT"], "STUDENT")
    mod_headers = _headers(seeded_users["MOD"], "MOD")
    registered = await client.post(
        f"/api/student/events/{standard_event}/register",
        json={"phoneNumber": "0901234567", "confirmCorrect": True},
        headers=student_headers,
    )
    assert registered.status_code == 201
    manual = await client.post(
        f"/api/mod/events/{standard_event}/attendance/check-in",
        json={"studentIds": [str(seeded_users["STUDENT"])]},
        headers=mod_headers,
    )
    assert manual.status_code == 200

    event_report = await client.get(
        f"/api/mod/events/{standard_event}/report/export", headers=mod_headers
    )
    assert event_report.status_code == 200
    workbook = load_workbook(BytesIO(event_report.content))
    assert workbook.active is not None
    assert workbook.active.freeze_panes == "A2"
    assert workbook.active.auto_filter.ref

    matrix = await client.get("/api/mod/reports/master-matrix", headers=mod_headers)
    assert matrix.status_code == 200
    assert matrix.json()["pagination"]["total"] >= 1
    matrix_export = await client.get("/api/mod/reports/master-matrix/export", headers=mod_headers)
    assert matrix_export.status_code == 200
    matrix_workbook = load_workbook(BytesIO(matrix_export.content))
    assert matrix_workbook.active is not None
    assert matrix_workbook.active.freeze_panes == "A2"
