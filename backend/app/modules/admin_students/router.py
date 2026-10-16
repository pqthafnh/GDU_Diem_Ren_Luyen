"""HTTP adapters for administrator student directory management."""

from io import BytesIO
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from openpyxl import load_workbook
from psycopg.errors import UniqueViolation

from app.modules.admin_students.schemas import (
    DeleteResult,
    StaffCreate,
    StaffDirectoryResponse,
    StaffResponse,
    StudentCreate,
    StudentDirectoryResponse,
    StudentResponse,
)
from app.modules.admin_students.service import AdminStudentService
from app.shared.auth import current_claims

router = APIRouter(tags=["admin-students"])
_HEADERS = {
    "studentId",
    "fullName",
    "email",
    "initialPassword",
    "gender",
    "birthDate",
    "className",
    "major",
    "phoneNumber",
    "enrollmentYear",
}
_STAFF_HEADERS = {
    "fullName",
    "email",
    "initialPassword",
    "role",
    "gender",
    "birthDate",
    "phoneNumber",
}


def _require_admin(claims: dict[str, object]) -> None:
    """Reject every role except ADMIN."""
    if claims.get("role") != "ADMIN":
        raise HTTPException(status_code=403, detail="ADMIN role required")


def _excel_students(content: bytes) -> list[StudentCreate]:
    """Parse the first worksheet using the documented camelCase header contract."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active
    if worksheet is None:
        raise ValueError("Excel file has no active worksheet")
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise ValueError("Excel file is empty")
    headers = [str(value).strip() if value is not None else "" for value in header_row]
    if not _HEADERS.issubset(headers):
        missing = sorted(_HEADERS.difference(headers))
        raise ValueError(f"Missing Excel columns: {', '.join(missing)}")
    result: list[StudentCreate] = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        data = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
        }
        try:
            result.append(StudentCreate.model_validate(data))
        except ValueError as error:
            raise ValueError(f"Invalid student data at Excel row {row_number}: {error}") from error
    return result


def _excel_staff(content: bytes) -> list[StaffCreate]:
    """Parse the first worksheet using the staff account header contract."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook.active
    if worksheet is None:
        raise ValueError("Excel file has no active worksheet")
    rows = worksheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        raise ValueError("Excel file is empty")
    headers = [str(value).strip() if value is not None else "" for value in header_row]
    if not _STAFF_HEADERS.issubset(headers):
        raise ValueError(
            f"Missing Excel columns: {', '.join(sorted(_STAFF_HEADERS.difference(headers)))}"
        )
    result: list[StaffCreate] = []
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        data = {
            headers[index]: values[index] if index < len(values) else None
            for index in range(len(headers))
        }
        try:
            result.append(StaffCreate.model_validate(data))
        except ValueError as error:
            raise ValueError(f"Invalid staff data at Excel row {row_number}: {error}") from error
    return result


@router.get("/students", response_model=StudentDirectoryResponse)
async def list_students(
    request: Request,
    search: str | None = Query(default=None, max_length=100),
    enrollment_year: int | None = Query(default=None, alias="enrollmentYear", ge=2000, le=2100),
    page: int = Query(default=1, ge=1),
    limit: int = Query(default=50, ge=1, le=200),
    claims: dict[str, object] = Depends(current_claims),
) -> StudentDirectoryResponse:
    """List student accounts for the authenticated administrator."""
    _require_admin(claims)
    async with request.app.state.database.connection() as connection:
        return await AdminStudentService().list_students(
            connection, search, enrollment_year, page, limit
        )


@router.post("/students", response_model=StudentResponse, status_code=201)
async def create_student(
    request: Request,
    body: StudentCreate,
    claims: dict[str, object] = Depends(current_claims),
) -> StudentResponse:
    """Create one student account for the authenticated administrator."""
    _require_admin(claims)
    try:
        async with request.app.state.database.connection() as connection:
            async with connection.transaction():
                return await AdminStudentService().create_student(connection, body)
    except UniqueViolation as error:
        raise HTTPException(status_code=409, detail="Student ID or email already exists") from error


@router.post("/students/import", response_model=list[StudentResponse], status_code=201)
async def import_students(
    request: Request,
    file: UploadFile = File(...),
    claims: dict[str, object] = Depends(current_claims),
) -> list[StudentResponse]:
    """Import student accounts from an XLSX workbook in one transaction."""
    _require_admin(claims)
    if file.content_type not in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    }:
        raise HTTPException(status_code=415, detail="Only .xlsx files are supported")
    try:
        content = await file.read()
        students = _excel_students(content)
        async with request.app.state.database.connection() as connection:
            return await AdminStudentService().create_students(connection, students)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    except UniqueViolation as error:
        raise HTTPException(status_code=409, detail="Student ID or email already exists") from error


@router.delete("/students/{student_id}", response_model=DeleteResult)
async def delete_student(
    request: Request,
    student_id: UUID,
    claims: dict[str, object] = Depends(current_claims),
) -> DeleteResult:
    """Delete one student account and its dependent records."""
    _require_admin(claims)
    async with request.app.state.database.connection() as connection:
        return await AdminStudentService().delete_student(connection, student_id)


@router.delete("/students/year/{enrollment_year}", response_model=DeleteResult)
async def delete_year(
    request: Request,
    enrollment_year: int,
    claims: dict[str, object] = Depends(current_claims),
) -> DeleteResult:
    """Delete all student accounts belonging to one enrollment year."""
    _require_admin(claims)
    async with request.app.state.database.connection() as connection:
        return await AdminStudentService().delete_year(connection, enrollment_year)


@router.get("/staff", response_model=StaffDirectoryResponse)
async def list_staff(
    request: Request,
    search: str | None = Query(default=None, max_length=100),
    role: str | None = Query(default=None, pattern=r"^(ADMIN|MOD)$"),
    page: int = Query(default=1, ge=1),
    limit: int = Query(default=50, ge=1, le=200),
    claims: dict[str, object] = Depends(current_claims),
) -> StaffDirectoryResponse:
    """List ADMIN and MOD accounts for the authenticated administrator."""
    _require_admin(claims)
    async with request.app.state.database.connection() as connection:
        return await AdminStudentService().list_staff(connection, search, role, page, limit)


@router.post("/staff", response_model=StaffResponse, status_code=201)
async def create_staff(
    request: Request,
    body: StaffCreate,
    claims: dict[str, object] = Depends(current_claims),
) -> StaffResponse:
    """Create one ADMIN or MOD account."""
    _require_admin(claims)
    try:
        async with request.app.state.database.connection() as connection:
            async with connection.transaction():
                return await AdminStudentService().create_staff(connection, body)
    except UniqueViolation as error:
        raise HTTPException(status_code=409, detail="Email already exists") from error


@router.post("/staff/import", response_model=list[StaffResponse], status_code=201)
async def import_staff(
    request: Request,
    file: UploadFile = File(...),
    claims: dict[str, object] = Depends(current_claims),
) -> list[StaffResponse]:
    """Import ADMIN and MOD accounts from an XLSX workbook atomically."""
    _require_admin(claims)
    if file.content_type not in {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    }:
        raise HTTPException(status_code=415, detail="Only .xlsx files are supported")
    try:
        staff = _excel_staff(await file.read())
        async with request.app.state.database.connection() as connection:
            return await AdminStudentService().create_staff_batch(connection, staff)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    except UniqueViolation as error:
        raise HTTPException(status_code=409, detail="Email already exists") from error


@router.delete("/staff/{staff_id}", response_model=DeleteResult)
async def delete_staff(
    request: Request,
    staff_id: UUID,
    claims: dict[str, object] = Depends(current_claims),
) -> DeleteResult:
    """Delete one ADMIN/MOD account without allowing self-deletion."""
    _require_admin(claims)
    try:
        actor_id = UUID(str(claims["userId"]))
        async with request.app.state.database.connection() as connection:
            return await AdminStudentService().delete_staff(connection, staff_id, actor_id)
    except (KeyError, ValueError) as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
