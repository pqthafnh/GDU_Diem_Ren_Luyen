"""Master matrix report queries and workbook generation."""

from io import BytesIO
from uuid import UUID

import anyio
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from psycopg import AsyncConnection

from app.modules.reports.service import sanitize_excel_value
from app.modules.training_points.service import classify_points


class MasterMatrixService:
    """Builds one row per student and one dynamic column per activity."""

    async def build(
        self,
        connection: AsyncConnection,
        actor_id: UUID,
        academic_year: str | None,
        semester: str | None,
        major: str | None,
        search: str | None,
        page: int,
        limit: int,
    ) -> dict[str, object]:
        """Return filtered, paginated matrix rows and activity metadata."""
        offset = (page - 1) * limit
        async with connection.cursor() as cursor:
            await cursor.execute("SELECT id, title FROM events ORDER BY start_time, id")
            activities = await cursor.fetchall()
            activity_ids = [row[0] for row in activities]
            await cursor.execute(
                """SELECT u.id, u.student_id, u.full_name, u.class_name, u.major,
                          COALESCE(SUM(l.points_awarded), 0)
                   FROM users u
                   LEFT JOIN training_points_ledger l ON l.student_id = u.id
                   WHERE u.role = 'STUDENT'
                     AND (%s::text IS NULL OR %s::text = 'all' OR u.major = %s)
                     AND (%s::text IS NULL OR u.student_id ILIKE %s OR u.full_name ILIKE %s
                          OR u.email ILIKE %s)
                   GROUP BY u.id
                   ORDER BY u.student_id
                   LIMIT %s OFFSET %s""",
                (
                    major,
                    major,
                    major,
                    search,
                    f"%{search}%" if search else None,
                    f"%{search}%" if search else None,
                    f"%{search}%" if search else None,
                    limit,
                    offset,
                ),
            )
            students = await cursor.fetchall()
            await cursor.execute(
                """SELECT student_id, event_id, COALESCE(SUM(points_awarded), 0)
                   FROM training_points_ledger
                   WHERE (%s::text IS NULL OR %s::text = '3-years-continuous' OR academic_year = %s)
                     AND (%s::text IS NULL OR %s::text = 'all' OR semester = %s)
                   GROUP BY student_id, event_id""",
                (academic_year, academic_year, academic_year, semester, semester, semester),
            )
            points = {(row[0], row[1]): int(row[2]) for row in await cursor.fetchall()}
            await cursor.execute(
                """SELECT COUNT(*) FROM users
                   WHERE role = 'STUDENT'
                     AND (%s::text IS NULL OR %s::text = 'all' OR major = %s)
                     AND (%s::text IS NULL OR student_id ILIKE %s OR full_name ILIKE %s
                          OR email ILIKE %s)""",
                (
                    major,
                    major,
                    major,
                    search,
                    f"%{search}%" if search else None,
                    f"%{search}%" if search else None,
                    f"%{search}%" if search else None,
                ),
            )
            count_row = await cursor.fetchone()
        rows: list[dict[str, object]] = []
        for student in students:
            values = [points.get((student[0], event_id), 0) for event_id in activity_ids]
            total = sum(values)
            rows.append(
                {
                    "studentId": student[0],
                    "studentCode": student[1],
                    "fullName": student[2],
                    "className": student[3],
                    "major": student[4],
                    "dynamicActivities": dict(
                        zip((event[1] for event in activities), values, strict=True)
                    ),
                    "totalPoints": total,
                    "classification": classify_points(total),
                }
            )
        return {
            "activities": [{"id": row[0], "title": row[1]} for row in activities],
            "rows": rows,
            "pagination": {
                "page": page,
                "limit": limit,
                "total": int(count_row[0]) if count_row else 0,
            },
        }

    async def workbook(self, data: dict[str, object]) -> BytesIO:
        """Generate a master matrix workbook outside the event loop."""
        return await anyio.to_thread.run_sync(self._build_workbook, data)

    @staticmethod
    def _build_workbook(data: dict[str, object]) -> BytesIO:
        workbook = Workbook()
        sheet = workbook.active
        if sheet is None:
            raise RuntimeError("Workbook did not create an active sheet")
        sheet.title = "Ma tran diem"
        activities = data["activities"]
        rows = data["rows"]
        if not isinstance(activities, list) or not isinstance(rows, list):
            raise RuntimeError("Invalid master matrix data")
        headers = (
            ["MSSV", "Họ tên", "Lớp", "Ngành"]
            + [str(item["title"]) for item in activities]
            + ["Tổng điểm", "Xếp loại"]
        )
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True)
        for row in rows:
            if not isinstance(row, dict):
                continue
            activity_values = row["dynamicActivities"]
            if not isinstance(activity_values, dict):
                continue
            sheet.append(
                [
                    sanitize_excel_value(row["studentCode"]),
                    sanitize_excel_value(row["fullName"]),
                    sanitize_excel_value(row["className"]),
                    sanitize_excel_value(row["major"]),
                    *[activity_values[str(item["title"])] for item in activities],
                    row["totalPoints"],
                    row["classification"],
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 18
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output
